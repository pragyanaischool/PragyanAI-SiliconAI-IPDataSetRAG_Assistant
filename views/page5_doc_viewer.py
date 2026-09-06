import streamlit as st
import os
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
import pandas as pd
from langchain_core.documents import Document
from langchain_groq import ChatGroq
from utils.database import db_get_all_groups_and_ips, db_get_file_registry

def render():
    st.image("PragyanAI_Transperent.png")
    st.title("📖 Page 5: Document Viewer & Page-by-Page RAG Summarizer")
    st.markdown("Select an IP group and model, browse ingested specification documents, view content page-by-page or slide-by-slide, and run AI RAG summarization on specific pages.")

    # Synchronize groups from SQLite if session state is empty
    if "ip_groups" not in st.session_state or not st.session_state.ip_groups:
        st.session_state.ip_groups = db_get_all_groups_and_ips()

    ip_groups = st.session_state.ip_groups
    if not ip_groups:
        st.warning("⚠️ No IP models found. Please ingest specification sheets or files on **Ingestion & IP Management** first.")
        return

    available_ips = list(ip_groups.keys())
    unique_groups = sorted(list(set(ip_groups.values())))

    # Step 1: Select Group and IP Model
    st.markdown("---")
    col_sel1, col_sel2 = st.columns(2)
    with col_sel1:
        selected_group = st.selectbox("Select Protocol Group / Family", unique_groups, key="viewer_group_sel")
        ips_in_group = [ip for ip in available_ips if ip_groups.get(ip, "General") == selected_group]
    with col_sel2:
        selected_ip = st.selectbox("Select IP Model / Version", ips_in_group, key="viewer_ip_sel")

    # Step 2: Select Document from SQLite Registry
    registry = db_get_file_registry(selected_ip)
    if not registry:
        st.info(f"💡 No registered files found under `{selected_ip}` in the database.")
        return

    source_names = list(registry.keys())
    selected_source = st.selectbox("Select Document File to View", source_names, key="viewer_source_sel")

    file_meta = registry[selected_source]
    doc_type = file_meta.get("type", "Document")
    total_pages = file_meta.get("pages", 1)

    st.markdown(f"**Document Type:** `{doc_type}` | **Total Pages / Sections / Slides:** `{total_pages}`")

    # Ensure raw docs list exists in session state for this IP
    if "ip_raw_docs" not in st.session_state:
        st.session_state.ip_raw_docs = {}
    if selected_ip not in st.session_state.ip_raw_docs:
        st.session_state.ip_raw_docs[selected_ip] = []

    # Check if raw chunks for this source are loaded in memory; if not, rebuild them from the file on disk
    source_docs = [
        d for d in st.session_state.ip_raw_docs.get(selected_ip, []) 
        if d.metadata.get("source") == selected_source
    ]

    if not source_docs:
        local_file_path = os.path.join("temp_ip_data", selected_source)
        # Handle cases where source might be a URL or stored path
        if not os.path.exists(local_file_path):
            # Try extracting filename from URL if it's a web spec
            base_name = selected_source.split("/")[-1].split("?")[0]
            local_file_path = os.path.join("temp_ip_data", base_name)

        if os.path.exists(local_file_path):
            with st.spinner(f"Loading document chunks into memory for `{selected_source}`..."):
                rebuilt_chunks = []
                if doc_type in ["PDF Datasheet", "Remote PDF Spec"] or local_file_path.lower().endswith(".pdf"):
                    with fitz.open(local_file_path) as pdf:
                        for p_idx, page in enumerate(pdf):
                            text = page.get_text()
                            if text.strip():
                                rebuilt_chunks.append(Document(
                                    page_content=text,
                                    metadata={"source": selected_source, "page": p_idx + 1, "ip": selected_ip, "group": selected_group, "type": doc_type}
                                ))
                elif doc_type == "Word Document" or local_file_path.lower().endswith(".docx"):
                    doc_obj = DocxDocument(local_file_path)
                    text = "\n".join([p.text for p in doc_obj.paragraphs if p.text.strip()])
                    rebuilt_chunks.append(Document(
                        page_content=text,
                        metadata={"source": selected_source, "page": 1, "ip": selected_ip, "group": selected_group, "type": doc_type}
                    ))
                elif doc_type == "Presentation Slide" or local_file_path.lower().endswith(".pptx"):
                    prs = Presentation(local_file_path)
                    for s_idx, slide in enumerate(prs.slides):
                        slide_text = ""
                        for shape in slide.shapes:
                            if shape.has_text_frame:
                                for p in shape.text_frame.paragraphs:
                                    slide_text += p.text + "\n"
                        if slide_text.strip():
                            rebuilt_chunks.append(Document(
                                page_content=slide_text,
                                metadata={"source": selected_source, "page": s_idx + 1, "ip": selected_ip, "group": selected_group, "type": doc_type}
                            ))
                elif doc_type == "Excel Register Map" or local_file_path.lower().endswith(".xlsx"):
                    xls = pd.ExcelFile(local_file_path)
                    for sheet in xls.sheet_names:
                        df = pd.read_excel(local_file_path, sheet_name=sheet)
                        sheet_md = f"Sheet: {sheet}\n\n" + df.to_markdown(index=False)
                        rebuilt_chunks.append(Document(
                            page_content=sheet_md,
                            metadata={"source": selected_source, "page": 1, "ip": selected_ip, "group": selected_group, "type": doc_type}
                        ))
                else:
                    with open(local_file_path, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read()
                    rebuilt_chunks.append(Document(
                        page_content=text,
                        metadata={"source": selected_source, "page": 1, "ip": selected_ip, "group": selected_group, "type": doc_type}
                    ))

                st.session_state.ip_raw_docs[selected_ip].extend(rebuilt_chunks)
                source_docs = rebuilt_chunks
        else:
            st.warning(f"⚠️ Local file cache for `{selected_source}` could not be found under `temp_ip_data/`. Please re-upload the file on Page 1 if necessary.")
            return

    # Step 3: Page Navigation Slider / Number Input
    st.markdown("---")
    col_nav1, col_nav2 = st.columns([3, 1])
    with col_nav1:
        page_num = st.slider("Navigate Page / Slide / Section", min_value=1, max_value=max(1, total_pages), value=1, key="doc_page_slider")
    with col_nav2:
        direct_page = st.number_input("Jump to Page", min_value=1, max_value=max(1, total_pages), value=page_num, key="doc_page_num_input")
        if direct_page != page_num:
            page_num = direct_page

    # Find the corresponding document chunk for the selected page
    current_chunk = next((d for d in source_docs if d.metadata.get("page", 1) == page_num), None)
    
    if not current_chunk and source_docs:
        idx = min(page_num - 1, len(source_docs) - 1)
        current_chunk = source_docs[idx]

    # Step 4: Display Content Depending on Document Type
    st.markdown(f"### 📄 Viewing: `{selected_source}` — [Page / Slide {page_num} of {total_pages}]")

    if "PDF" in doc_type:
        st.info("📌 **PDF Page View Mode**")
        st.text_area("Page Text Content", value=current_chunk.page_content if current_chunk else "", height=350, key="pdf_text_view")
        
        # Optional: Render visual PDF page preview if local file exists
        local_pdf_path = os.path.join("temp_ip_data", selected_source)
        if not os.path.exists(local_pdf_path):
            local_pdf_path = os.path.join("temp_ip_data", selected_source.split("/")[-1].split("?")[0])

        if os.path.exists(local_pdf_path):
            try:
                with fitz.open(local_pdf_path) as pdf_file:
                    if page_num <= len(pdf_file):
                        page_obj = pdf_file[page_num - 1]
                        pix = page_obj.get_pixmap(dpi=150)
                        img_path = f"temp_ip_data/page_{page_num}.png"
                        pix.save(img_path)
                        st.image(img_path, caption=f"Rendered View - Page {page_num}", use_container_width=True)
            except Exception:
                pass

    elif doc_type == "Presentation Slide":
        st.info("🖥️ **Presentation Slide View Mode**")
        st.markdown(f"#### Slide {page_num}")
        st.code(current_chunk.page_content if current_chunk else "", language="markdown")

    elif doc_type == "Excel Register Map":
        st.info("📊 **Excel Register Map View Mode**")
        st.markdown(current_chunk.page_content if current_chunk else "")

    elif doc_type == "HDL Code":
        st.info("⚡ **RTL Source Code View Mode**")
        st.code(current_chunk.page_content if current_chunk else "", language="verilog")

    else:
        st.info("📝 **General Document View Mode**")
        st.text_area("Content", value=current_chunk.page_content if current_chunk else "", height=350, key="gen_text_view")

    # Step 5: Page-by-Page RAG Summarizer
    st.markdown("---")
    st.subheader("🤖 Page-by-Page AI RAG Summarizer & Technical Analysis")
    
    analysis_focus = st.selectbox(
        "Select Summarization Focus", 
        [
            "Comprehensive Technical Summary", 
            "Extract Register Definitions & Bitfields", 
            "Protocol Handshaking & State Machine Analysis", 
            "Timing, Constraints & Safety Rules"
        ],
        key="sum_focus_sel"
    )

    if st.button("✨ Summarize & Analyze Selected Page", type="primary", key="btn_summarize_page"):
        try:
            groq_api_key = st.secrets["GROQ_API_KEY"]
            model_name = st.secrets.get("MODEL_NAME", "llama-3.3-70b-versatile")
        except Exception:
            st.error("⚠️ `GROQ_API_KEY` or `MODEL_NAME` not found in `st.secrets`.")
            return

        llm = ChatGroq(model=model_name, temperature=0.1, groq_api_key=groq_api_key)

        page_text = current_chunk.page_content if current_chunk else "No content available."

        rag_summary_prompt = f"""You are a Principal Silicon Architect and Hardware Verification Expert. 
Analyze the following text extracted from **{selected_source}** (IP Model: `{selected_ip}`, Page/Slide: `{page_num}`).

### Analysis Focus:
{analysis_focus}

### Target Page Content:
{page_text}

### Instructions:
Provide a rigorous, engineering-grade breakdown of this specific page content, highlighting critical parameters, register offsets, signal names, or protocol rules.
"""

        with st.spinner(f"Running AI analysis on Page {page_num} using `{model_name}`..."):
            response = llm.invoke(rag_summary_prompt)
            st.markdown("### 📋 Page AI Analysis & Summary Result")
            st.markdown(response.content)
